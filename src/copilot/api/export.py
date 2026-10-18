"""Excel downloads of the statements and the 13-week forecast (web dashboard).

The workbooks are built from the same gateway results the dashboard shows (role-checked and
audited like every other read), already turned into JSON. Nothing is recomputed: each money
cell holds the API's own value, only written as a number so Excel can format it.

Imported text (business, counterparty and document names inside forecast labels, saved
assumptions) is untrusted. A spreadsheet would run a cell starting with "=" as a formula
(e.g. =HYPERLINK(...)), so every text cell is written as a string, and one that begins with a
formula character also gets Excel's quote prefix (known risk R53).
"""

from __future__ import annotations

from decimal import Decimal
from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.cell.cell import Cell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.worksheet import Worksheet

MONEY = "#,##0.00;[Red]-#,##0.00"
FORMULA_START = ("=", "+", "-", "@", "\t", "\r", "\n")
NOTE = "Figures come from your books as imported into Fintellect. Not audited."
_BOLD = Font(bold=True)
_HEAD = Font(bold=True, color="FFFFFF")
_HEAD_FILL = PatternFill("solid", fgColor="0F7A63")

# Line labels match the Statements page (js/pages/statements.js).
OPEX = {
    "payroll": "Payroll",
    "rent": "Rent",
    "utilities": "Utilities",
    "software": "Software",
    "other": "Other costs",
    "depreciation": "Depreciation",
}
WORKING = {
    "receivables": "Customer invoices unpaid",
    "inventory": "Stock",
    "payables": "Supplier bills unpaid",
    "gst_payable": "GST owed",
    "cpf_payable": "CPF owed",
    "unmatched_customer_receipts": "Customer payments not yet matched",
    "unmatched_supplier_payments": "Supplier payments not yet matched",
}
ASSETS = [
    ("cash", "Cash"),
    ("receivables", "Customer invoices unpaid"),
    ("inventory", "Stock"),
    ("unmatched_supplier_payments", "Supplier payments not yet matched"),
    ("fixed_assets_net", "Equipment and vehicles (after depreciation)"),
]
LIABILITIES = [
    ("payables", "Supplier bills unpaid"),
    ("gst_payable", "GST owed"),
    ("cpf_payable", "CPF owed"),
    ("unmatched_customer_receipts", "Customer payments not yet matched"),
    ("accrued_costs", "Accrued costs"),
    ("loans_due_within_12_months", "Loans due within 12 months"),
    ("loans_due_after_12_months", "Loans due after 12 months"),
    ("other_borrowings", "Other borrowings"),
]
EQUITY = [
    ("share_capital", "Share capital"),
    ("retained_earnings", "Retained earnings"),
    ("owner_drawings", "Owner drawings"),
]
PERIODS = {
    "month": "Month",
    "quarter": "Quarter",
    "ytd": "Year to date",
    "last_12_months": "Last 12 months",
}
VIEWS = {
    "as_imported": "As imported (no adjustments)",
    "with_adjustments": "With recorded adjustments",
}
MARGINS = [
    ("gross_margin", "Gross margin"),
    ("operating_margin", "Operating margin"),
    ("net_margin", "Net margin"),
]
CONTENT_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def text_cell(cell: Cell, text: Any) -> None:
    """Write text that is never treated as a formula (R53)."""
    value = "" if text is None else str(text)
    cell.value = value
    cell.data_type = "s"
    if value.startswith(FORMULA_START):
        cell.quotePrefix = True


def money_cell(cell: Cell, amount: Any) -> None:
    cell.value = Decimal(str(amount))
    cell.number_format = MONEY


class Sheet:
    """Appends rows: a label column, then money or text columns."""

    def __init__(self, ws: Worksheet, widths: list[int]):
        self.ws = ws
        self.row = 1
        for i, w in enumerate(widths):
            ws.column_dimensions[chr(ord("A") + i)].width = w

    def header(self, *titles: str) -> None:
        for i, t in enumerate(titles, start=1):
            c = self.ws.cell(self.row, i)
            text_cell(c, t)
            c.font, c.fill = _HEAD, _HEAD_FILL
            if i > 1:
                c.alignment = Alignment(horizontal="right")
        self.ws.freeze_panes = self.ws.cell(self.row + 1, 1)
        self.row += 1

    def line(self, label: str, *values: Any, bold: bool = False, kind: str = "money") -> None:
        """kind: "money" (S$ amounts), "pct" (a fraction shown as %), "num" or "text"."""
        c = self.ws.cell(self.row, 1)
        text_cell(c, label)
        if bold:
            c.font = _BOLD
        for i, v in enumerate(values, start=2):
            cell = self.ws.cell(self.row, i)
            if v is None:
                text_cell(cell, "Not available")
            elif kind == "text":
                text_cell(cell, v)
            else:
                cell.value = Decimal(str(v))
                cell.number_format = {"money": MONEY, "pct": "0.0%", "num": "0.00"}[kind]
            if bold:
                cell.font = _BOLD
        self.row += 1

    def group(self, label: str) -> None:
        self.row += 1 if self.row > 2 else 0
        self.line(label, bold=True)

    def blank(self) -> None:
        self.row += 1


def _about(wb: Workbook, title: str, rows: list[tuple[str, Any]]) -> None:
    sh = Sheet(wb.active, [34, 70])
    sh.ws.title = "About"
    c = sh.ws.cell(1, 1)
    text_cell(c, title)
    c.font = Font(bold=True, size=14)
    sh.row = 3
    for label, value in rows:
        sh.line(label, value, kind="text")
    sh.blank()
    sh.line(NOTE)


def _check(ok: bool, difference: Any) -> str:
    return f"{'Yes' if ok else 'No'} (difference {difference})"


def _bytes(wb: Workbook) -> bytes:
    out = BytesIO()
    wb.save(out)
    return out.getvalue()


def statements_workbook(business_name: str, v: dict[str, Any]) -> bytes:
    """v = the /api/statements JSON's data.value (all three statements)."""
    wb = Workbook()
    checks = v["checks"]
    _about(
        wb,
        f"Financial statements: {business_name}",
        [
            ("Business", business_name),
            ("Period", f"{PERIODS.get(v['period'], v['period'])}: {v['start']} to {v['end']}"),
            ("Basis", v["basis"]),
            ("View", VIEWS.get(v["view"], v["view"])),
            (
                "Balance sheet balances",
                _check(checks["balanced"], checks["balance_sheet_difference"]),
            ),
            (
                "Closing cash equals the bank",
                _check(checks["cash_matches_bank"], checks["closing_cash_minus_bank"]),
            ),
        ],
    )

    p = v["pnl"]
    sh = Sheet(wb.create_sheet("Profit and loss"), [44, 18])
    sh.header("", f"{v['start']} to {v['end']}")
    sh.line("Sales", p["revenue"])
    sh.line("Cost of sales", p["cost_of_sales"])
    sh.line("Gross profit", p["gross_profit"], bold=True)
    sh.group("Operating costs")
    for k, amount in p["operating_expenses"].items():
        sh.line(OPEX.get(k, k.replace("_", " ").capitalize()), amount)
    sh.line("Total operating costs", p["total_operating_expenses"], bold=True)
    sh.line("Operating profit", p["operating_profit"], bold=True)
    sh.line("Other income", p["other_income"])
    sh.line("Interest", p["interest"])
    sh.line("Net profit", p["net_profit"], bold=True)
    sh.blank()
    r = v["ratios"]
    for key, label in MARGINS:
        sh.line(label, r[key], kind="pct")
    sh.line("Current ratio", r["current_ratio"], kind="num")

    b = v["balance_sheet"]
    sh = Sheet(wb.create_sheet("Balance sheet"), [44, 18])
    sh.header("", f"At {b['as_of']}")
    sh.line("Assets", bold=True)
    for k, label in ASSETS:
        sh.line(label, b[k])
    sh.line("Total assets", b["total_assets"], bold=True)
    sh.group("Liabilities")
    for k, label in LIABILITIES:
        sh.line(label, b[k])
    sh.line("Total liabilities", b["total_liabilities"], bold=True)
    sh.group("Equity")
    for k, label in EQUITY:
        sh.line(label, b[k])
    sh.line("Total equity", b["total_equity"], bold=True)

    c = v["cash_flow"]
    sh = Sheet(wb.create_sheet("Cash flow"), [44, 18])
    sh.header("", f"{v['start']} to {v['end']}")
    sh.line("Net profit", c["net_profit"])
    sh.line("Add back depreciation", c["add_depreciation"])
    sh.group("Working capital")
    for k, amount in c["working_capital"].items():
        sh.line(WORKING.get(k, k.replace("_", " ").capitalize()), amount)
    sh.line("Cash from operations", c["operating"], bold=True)
    sh.line("Investing", c["investing"], bold=True)
    sh.line("Loan principal repaid", c["financing_loan_principal"])
    sh.line("New borrowing", c["new_borrowing"])
    sh.line("Owner drawings", c["owner_drawings"])
    sh.line("Cash from financing", c["financing"], bold=True)
    sh.line("Net change in cash", c["net_change"], bold=True)
    sh.line("Opening cash", c["opening_cash"])
    sh.line("Closing cash", c["closing_cash"], bold=True)
    return _bytes(wb)


def forecast_workbook(business_name: str, f: dict[str, Any]) -> bytes:
    """f = the forecast_balance result's value, as JSON (the Overview's forecast panel)."""
    wb = Workbook()
    low = f["low_point"]
    _about(
        wb,
        f"13-week cash forecast: {business_name}",
        [
            ("Business", business_name),
            ("Forecast as of", f["as_of"]),
            ("Starting cash", f["start_balance"]),
            ("Lowest point", f"{low['balance']} on {low['date']} (week {low['week']})"),
        ]
        + [("Assumption", a) for a in f["assumptions"]]
        + [("Saved assumption", a) for a in f.get("saved_assumptions") or []],
    )
    sh = Sheet(wb.create_sheet("Weekly"), [8, 13, 13, 16, 16, 16, 16])
    sh.header("Week", "Start", "End", "Money in", "Money out", "End balance", "Lowest in week")
    for w in f["weeks"]:
        row = sh.row
        sh.ws.cell(row, 1).value = int(w["week"])
        for col, key in ((2, "start"), (3, "end")):
            text_cell(sh.ws.cell(row, col), w[key])
        for col, key in ((4, "inflows"), (5, "outflows"), (6, "end_balance"), (7, "min_balance")):
            money_cell(sh.ws.cell(row, col), w[key])
        sh.row += 1

    sh = Sheet(wb.create_sheet("Items"), [13, 16, 14, 70])
    sh.header("Date", "Amount", "Source", "What")
    for it in f["items"]:
        row = sh.row
        text_cell(sh.ws.cell(row, 1), it["date"])
        money_cell(sh.ws.cell(row, 2), it["amount"])
        text_cell(sh.ws.cell(row, 3), it["source"])
        text_cell(sh.ws.cell(row, 4), it["label"])
        sh.row += 1
    return _bytes(wb)
