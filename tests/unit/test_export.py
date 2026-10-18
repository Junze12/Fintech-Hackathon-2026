"""Excel downloads of the statements and the forecast (web dashboard, 2026-10-17).

The workbooks must hold exactly the numbers the API returns (nothing recomputed), go through
the gateway like every read, and never let imported text become a spreadsheet formula (R53).
"""

from decimal import Decimal
from io import BytesIO

import pytest

pytest.importorskip("fastapi", reason="needs the api extra")
openpyxl = pytest.importorskip("openpyxl", reason="needs the api extra")

from fastapi.testclient import TestClient  # noqa: E402
from openpyxl import Workbook  # noqa: E402

from copilot.agent.mock_rules import RuleBasedLLM  # noqa: E402
from copilot.api import export  # noqa: E402
from copilot.api.app import create_app  # noqa: E402
from copilot.api.sessions import SessionStore  # noqa: E402
from copilot.runtime import open_scenario  # noqa: E402

STATEMENTS = "/api/export/statements.xlsx"
FORECAST = "/api/export/forecast.xlsx"


def signed_in(scenario: str = "cash_crunch", role: str = "owner") -> TestClient:
    c = TestClient(
        create_app(SessionStore(), opener=lambda f, r, a: open_scenario(f, r, llm=RuleBasedLLM()))
    )
    assert c.post("/api/session", json={"scenario": scenario, "role": role}).status_code == 200
    return c


def book(r):
    assert r.status_code == 200, r.text
    assert r.headers["content-type"] == export.CONTENT_TYPE
    assert r.headers["content-disposition"].startswith("attachment; filename=")
    assert r.headers["x-content-type-options"] == "nosniff"
    assert r.headers["cache-control"] == "no-store"
    return openpyxl.load_workbook(BytesIO(r.content))


def labelled(ws) -> dict[str, Decimal]:
    """Column A label -> column B number (first occurrence)."""
    out: dict[str, Decimal] = {}
    for a, b in ws.iter_rows(min_col=1, max_col=2, values_only=True):
        if a and isinstance(b, (int, float)) and a not in out:
            out[a] = Decimal(str(b)).quantize(Decimal("0.01"))
    return out


def d(x) -> Decimal:
    return Decimal(str(x)).quantize(Decimal("0.01"))


@pytest.mark.parametrize(
    ("query", "net"),
    [
        ({"period": "month", "month": "2026-09"}, "-4105.30"),
        ({"period": "last_12_months"}, "-22467.49"),
    ],
)
def test_statements_workbook_matches_the_api_cell_for_cell(query, net):
    c = signed_in()
    v = c.get("/api/statements", params=query).json()["data"]["value"]
    wb = book(c.get(STATEMENTS, params=query))
    assert wb.sheetnames == ["About", "Profit and loss", "Balance sheet", "Cash flow"]

    pnl = labelled(wb["Profit and loss"])
    p = v["pnl"]
    assert pnl["Net profit"] == d(p["net_profit"]) == Decimal(net)
    for label, key in [
        ("Sales", "revenue"),
        ("Cost of sales", "cost_of_sales"),
        ("Gross profit", "gross_profit"),
        ("Total operating costs", "total_operating_expenses"),
        ("Operating profit", "operating_profit"),
        ("Interest", "interest"),
    ]:
        assert pnl[label] == d(p[key]), label
    for key, amount in p["operating_expenses"].items():
        assert pnl[export.OPEX[key]] == d(amount), key

    bs = labelled(wb["Balance sheet"])
    b = v["balance_sheet"]
    for key, label in export.ASSETS + export.LIABILITIES + export.EQUITY:
        if label not in ("Customer payments not yet matched", "Supplier payments not yet matched"):
            assert bs[label] == d(b[key]), label
    for label, key in [
        ("Total assets", "total_assets"),
        ("Total liabilities", "total_liabilities"),
        ("Total equity", "total_equity"),
    ]:
        assert bs[label] == d(b[key]), label

    cf = labelled(wb["Cash flow"])
    f = v["cash_flow"]
    for label, key in [
        ("Cash from operations", "operating"),
        ("Cash from financing", "financing"),
        ("Net change in cash", "net_change"),
        ("Opening cash", "opening_cash"),
        ("Closing cash", "closing_cash"),
    ]:
        assert cf[label] == d(f[key]), label

    about = {a: b for a, b in wb["About"].iter_rows(min_col=1, max_col=2, values_only=True) if a}
    assert about["Balance sheet balances"] == "Yes (difference 0.00)"
    assert about["Closing cash equals the bank"] == "Yes (difference 0.00)"


def test_money_cells_are_numbers_with_a_money_format():
    wb = book(signed_in().get(STATEMENTS))
    cell = next(r[1] for r in wb["Profit and loss"].iter_rows() if r[0].value == "Net profit")
    assert isinstance(cell.value, (int, float)) and cell.number_format == export.MONEY


def test_forecast_workbook_matches_the_forecast_panel():
    c = signed_in()
    f = c.get("/api/dashboard").json()["panels"]["forecast"]["data"]["value"]
    wb = book(c.get(FORECAST))
    assert wb.sheetnames == ["About", "Weekly", "Items"]
    weeks = list(wb["Weekly"].iter_rows(min_row=2, values_only=True))
    assert len(weeks) == len(f["weeks"]) == 13
    for row, w in zip(weeks, f["weeks"], strict=True):
        assert row[0] == w["week"] and row[1] == w["start"] and row[2] == w["end"]
        assert [d(x) for x in row[3:7]] == [
            d(w["inflows"]),
            d(w["outflows"]),
            d(w["end_balance"]),
            d(w["min_balance"]),
        ]
    items = list(wb["Items"].iter_rows(min_row=2, values_only=True))
    assert [(r[0], d(r[1]), r[2], r[3]) for r in items] == [
        (i["date"], d(i["amount"]), i["source"], i["label"]) for i in f["items"]
    ]
    about = {a: b for a, b in wb["About"].iter_rows(min_col=1, max_col=2, values_only=True) if a}
    low = f["low_point"]
    assert about["Lowest point"] == f"{low['balance']} on {low['date']} (week {low['week']})"


@pytest.mark.parametrize("text", ['=HYPERLINK("http://x","click")', "+1", "-2", "@SUM(A1)", "\t=1"])
def test_text_that_looks_like_a_formula_stays_text(text):
    wb = Workbook()
    cell = wb.active["A1"]
    export.text_cell(cell, text)
    out = BytesIO()
    wb.save(out)
    back = openpyxl.load_workbook(BytesIO(out.getvalue())).active["A1"]
    assert back.value == text and back.data_type == "s" and back.quotePrefix


def test_poisoned_names_come_out_as_plain_text():
    c = signed_in("prompt_injection")
    name = c.get("/api/session").json()["business_name"]
    for path in (STATEMENTS, FORECAST):
        wb = book(c.get(path))
        cells = [x for ws in wb for row in ws.iter_rows() for x in row if x.value is not None]
        assert not [x for x in cells if x.data_type == "f"], path  # no formulas anywhere
        assert any(name in str(x.value) for x in cells), path  # shown literally, unchanged


def test_downloads_follow_the_session_and_the_gateway():
    anon = TestClient(create_app(SessionStore()))
    assert anon.get(STATEMENTS).status_code == 401
    assert anon.get(FORECAST).status_code == 401
    for role in ("owner", "staff", "viewer"):
        c = signed_in(role=role)
        for path, api in ((STATEMENTS, "/api/statements"), (FORECAST, "/api/dashboard")):
            r = c.get(path)
            ok = (
                c.get(api).json()["ok"]
                if api == "/api/statements"
                else c.get(api).json()["panels"]["forecast"]["ok"]
            )
            assert r.status_code == (200 if ok else 403), (role, path)


def test_bad_parameters_are_rejected_like_the_statements_endpoint():
    c = signed_in()
    assert c.get(STATEMENTS, params={"period": "decade"}).status_code == 422
    assert c.get(STATEMENTS, params={"month": "2026-13"}).status_code == 422
    assert c.get(STATEMENTS, params={"view": "<script>"}).status_code == 422


def test_file_names_hold_only_safe_characters():
    c = signed_in()
    r = c.get(STATEMENTS, params={"period": "quarter"})
    name = r.headers["content-disposition"].split('filename="')[1].rstrip('"')
    assert name == "biz-straits-statements-quarter-2026-09-30.xlsx"
    assert (
        c.get(FORECAST)
        .headers["content-disposition"]
        .endswith('filename="biz-straits-forecast-2026-09-30.xlsx"')
    )
